"""Tests de la exportación a Excel."""
from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook


def abrir(respuesta):
    assert respuesta.status_code == 200, respuesta.text
    assert respuesta.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in respuesta.headers["content-disposition"]
    return load_workbook(BytesIO(respuesta.content))


def test_exportar_lo_que_se_ve(cliente):
    r = cliente.post("/exportar", json={
        "nombre": "stock insumos", "titulo": "Stock",
        "columnas": ["Insumo", "Cantidad", "Fecha", "Hora"],
        "filas": [["Urea", 1500.5, "2026-03-15", "2026-03-15 10:30:00"], ["Gasoil", 20, None, None]],
    })
    assert 'filename="stock_insumos_' in r.headers["content-disposition"]
    hoja = abrir(r)["Stock"]
    assert [c.value for c in hoja[1]] == ["Insumo", "Cantidad", "Fecha", "Hora"]
    assert hoja["B2"].value == 1500.5
    assert hoja["C2"].value == datetime(2026, 3, 15)  # Fecha de verdad, no texto.
    assert hoja["C2"].number_format == "DD/MM/YYYY"
    assert hoja.freeze_panes == "A2"


def test_exportar_todo(cliente, bot):
    bot("/nuevo glifosato herbicida litros")
    bot("/entrada 20 glifosato - compra")
    bot("/nuevo urea fertilizante kg")
    bot("/nuevo filtro filtros unidades")
    cliente.post("/maquinas", json={"nombre": "JD", "tipo": "tractor"})
    cliente.post("/animales", json={"caravana": "1234", "categoria_id": 1})  # 1 = Vaca (Vacuno)
    libro = abrir(cliente.get("/exportar/completo"))
    assert libro.sheetnames == ["Insumos", "Químicos", "Repuestos", "Movimientos", "Máquinas", "Services y arreglos",
                                "Service programado", "Trabajos", "Vencimientos", "Contactos", "Animales",
                                "Eventos de animales", "Lotes", "Cultivos por lote"]
    # Cada insumo en su hoja, igual que en la web.
    assert [c.value for c in libro["Químicos"][2]][:4] == ["glifosato", "Agroquímico", "Herbicida", 20]
    assert [fila[0].value for fila in libro["Insumos"].iter_rows(min_row=2)] == ["urea"]
    assert [fila[0].value for fila in libro["Repuestos"].iter_rows(min_row=2)] == ["filtro"]
    assert libro["Movimientos"]["F2"].value == "compra"
    assert libro["Animales"]["A2"].value == "1234"


def test_exportar_rechaza_datos_raros(cliente):
    r = cliente.post("/exportar", json={"columnas": [], "filas": []})
    assert r.status_code == 422
