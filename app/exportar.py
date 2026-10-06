"""Exportar datos a Excel (.xlsx) con la librería openpyxl.

Dos formas:
- POST /exportar            -> arma un Excel con las filas que manda la web
                               (exactamente lo que estás viendo, con los filtros aplicados).
- GET  /exportar/completo   -> un Excel con TODAS las tablas, una hoja por tabla.
"""
import re
from datetime import date, datetime
from io import BytesIO
from typing import Optional, Union

from fastapi import APIRouter
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from app.ganaderia import db as ganaderia_db
from app.insumos import db as insumos_db
from app.lotes import db as lotes_db
from app.maquinaria import db as maquinaria_db
from app.nucleo import opciones

router = APIRouter(tags=["Exportar"])

TIPO_EXCEL = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FONDO_TITULOS = PatternFill("solid", fgColor="E7F6F0")  # Verde muy claro (el acento de la web).


# ---------- Armar el archivo ----------

def _convertir(valor):
    """Las fechas llegan como texto: las pasamos a fecha de verdad para que Excel las entienda."""
    if isinstance(valor, str):
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", valor):
            return date.fromisoformat(valor)
        if re.fullmatch(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}(:\d{2})?", valor):
            return datetime.fromisoformat(valor)
    return valor


def _agregar_hoja(libro, titulo: str, columnas: list, filas: list):
    hoja = libro.create_sheet(titulo[:31])  # Excel no acepta nombres de hoja de más de 31 letras.
    hoja.append(columnas)
    for celda in hoja[1]:
        celda.font = Font(bold=True)
        celda.fill = FONDO_TITULOS
        celda.alignment = Alignment(vertical="center")

    anchos = [len(str(c)) for c in columnas]
    for fila in filas:
        valores = [_convertir(v) for v in fila]
        hoja.append(valores)
        for i, valor in enumerate(valores):
            anchos[i] = max(anchos[i], len(str(valor)) if valor is not None else 0)

    # Formatos: fechas dd/mm/aaaa y números con separador de miles.
    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            if isinstance(celda.value, datetime):
                celda.number_format = "DD/MM/YYYY HH:MM"
            elif isinstance(celda.value, date):
                celda.number_format = "DD/MM/YYYY"
            elif isinstance(celda.value, float) and not celda.value.is_integer():
                celda.number_format = "#,##0.00"
            elif isinstance(celda.value, (int, float)) and not isinstance(celda.value, bool):
                celda.number_format = "#,##0"

    for i, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = min(max(ancho + 2, 10), 60)
    hoja.freeze_panes = "A2"  # La fila de títulos queda fija al bajar.
    if filas:
        hoja.auto_filter.ref = hoja.dimensions  # Flechitas de filtro en los títulos.


def crear_excel(hojas: list) -> bytes:
    """hojas = [(titulo, columnas, filas), ...]  ->  el archivo .xlsx en bytes."""
    libro = Workbook()
    libro.remove(libro.active)  # Sacamos la hoja vacía que viene por defecto.
    for titulo, columnas, filas in hojas:
        _agregar_hoja(libro, titulo, columnas, filas)
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _descarga(contenido: bytes, nombre: str) -> Response:
    """Respuesta que el navegador guarda como archivo (en vez de mostrarla)."""
    nombre = re.sub(r"[^\w\-]+", "_", nombre).strip("_") or "agroapp"
    return Response(
        content=contenido,
        media_type=TIPO_EXCEL,
        headers={"Content-Disposition": f'attachment; filename="{nombre}_{date.today():%Y-%m-%d}.xlsx"'},
    )


# ---------- Lo que se ve en pantalla ----------

Valor = Union[str, float, int, bool, None]


class PedidoExportar(BaseModel):
    nombre: str = Field(default="agroapp", max_length=60)  # Nombre del archivo, sin extensión.
    titulo: str = Field(default="Datos", max_length=31)    # Nombre de la hoja.
    columnas: list[str] = Field(min_length=1, max_length=40)
    filas: list[list[Valor]] = Field(max_length=50000)


@router.post("/exportar")
def exportar_tabla(pedido: PedidoExportar):
    return _descarga(crear_excel([(pedido.titulo, pedido.columnas, pedido.filas)]), pedido.nombre)


# ---------- Todo junto ----------

def _etiqueta(diccionario, valor):
    return diccionario.get(valor, valor or "")


def _subcategoria(insumo):
    return opciones.subcategorias_de(insumo["categoria"]).get(insumo["subcategoria"], "")


def hojas_completas() -> list:
    """Una hoja por tabla, con nombres entendibles (no los códigos internos)."""
    insumos = insumos_db.listar_insumos(incluir_archivados=True)
    maquinas = maquinaria_db.listar_maquinas(incluir_archivadas=True)
    nombre_maquina = {m["id"]: m["nombre"] for m in maquinas}
    animales = ganaderia_db.listar_animales(incluir_bajas=True)
    caravana = {a["id"]: a["caravana"] for a in animales}
    especie = {a["id"]: a["especie"] for a in animales}

    mantenimientos, trabajos, eventos = [], [], []
    for m in maquinas:
        mantenimientos += maquinaria_db.listar_mantenimientos(m["id"])
        trabajos += maquinaria_db.listar_trabajos(maquina_id=m["id"])
    for a in animales:
        eventos += ganaderia_db.listar_eventos(a["id"])

    C, E = opciones.CATEGORIAS, opciones.ESTADOS_REPRODUCTIVOS

    def hoja_insumos(hoja):
        """Insumos, Químicos y Repuestos van en hojas separadas (como en la web)."""
        return (opciones.HOJAS_INSUMOS[hoja],
                ["Nombre", "Categoría", "Tipo", "Cantidad", "Unidad", "Stock mínimo", "Máquinas", "Archivado"],
                [[i["nombre"], _etiqueta(C, i["categoria"]), _subcategoria(i), i["cantidad"], i["unidad"],
                  i["stock_minimo"], ", ".join(m["nombre"] for m in i["maquinas"]),
                  "Sí" if i["archivado"] else "No"] for i in insumos if i["hoja"] == hoja])

    return [
        *(hoja_insumos(hoja) for hoja in opciones.HOJAS_INSUMOS),
        ("Movimientos", ["Fecha", "Insumo", "Tipo", "Cantidad", "Unidad", "Motivo"],
         [[m["fecha"], m["insumo_nombre"], m["tipo"].capitalize(), m["cantidad"], m["unidad"], m["motivo"]]
          for m in insumos_db.buscar_movimientos(limite=1_000_000)]),
        ("Máquinas", ["Nombre", "Tipo", "Marca", "Modelo", "Año", "N° de serie", "N° serie monitor", "Patente", "Horas motor",
                      "Horas trilla", "Hectáreas trabajadas", "Archivada", "Observaciones"],
         [[m["nombre"], _etiqueta(opciones.TIPOS_MAQUINA, m["tipo"]), m["marca"], m["modelo"], m["anio"],
           m["numero_serie"], m["serie_monitor"], m["patente"], m["horas_motor"], m["horas_trilla"], m["hectareas_totales"],
           "Sí" if m["archivado"] else "No", m["observaciones"]] for m in maquinas]),
        ("Services y arreglos", ["Fecha", "Máquina", "Tipo", "Horas", "Descripción", "Costo"],
         [[s["fecha"], nombre_maquina.get(s["maquina_id"]), _etiqueta(opciones.TIPOS_MANTENIMIENTO, s["tipo"]),
           s["horas"], s["descripcion"], s["costo"]] for s in sorted(mantenimientos, key=lambda x: x["fecha"], reverse=True)]),
        ("Service programado", ["Máquina", "Plan", "Cada (h)", "Medida", "Último (h)", "Próximo (h)", "Faltan (h)",
                                "Estado", "Activo"],
         [[p["maquina_nombre"], p["nombre"], p["cada_horas"], _etiqueta(opciones.MEDIDAS_SERVICE, p["medida"]),
           p["ultima_horas"], p["proximo"], p["faltan"], {"ok": "Al día", "proximo": "Próximo", "vencido": "Vencido"}[p["estado"]],
           "Sí" if p["activo"] else "No"] for p in maquinaria_db.listar_planes(solo_activos=False)]),
        ("Trabajos", ["Fecha", "Máquina", "Trabajo", "Hectáreas", "Lote", "Cultivo", "Observaciones"],
         [[t["fecha"], t["maquina_nombre"], _etiqueta(opciones.TIPOS_TRABAJO, t["tipo"]), t["hectareas"], t["lote"],
           t["cultivo"], t["observaciones"]] for t in sorted(trabajos, key=lambda x: x["fecha"], reverse=True)]),
        ("Vencimientos", ["Vence", "Descripción", "Tipo", "Máquina", "N° serie monitor", "Resuelto", "Observaciones"],
         [[v["fecha_vencimiento"], v["descripcion"], _etiqueta(opciones.TIPOS_VENCIMIENTO, v["tipo"]),
           v["maquina_nombre"] or "", (v["maquina_serie_monitor"] or "") if v["tipo"] in opciones.TIPOS_CON_MONITOR else "", "Sí" if v["resuelto"] else "No", v["observaciones"]]
          for v in maquinaria_db.listar_vencimientos(incluir_resueltos=True)]),
        ("Contactos", ["Nombre", "Rubro", "Empresa", "Teléfono", "Email", "Notas"],
         [[c["nombre"], _etiqueta(opciones.RUBROS_CONTACTO, c["rubro"]), c["empresa"], c["telefono"], c["email"],
           c["notas"]] for c in maquinaria_db.listar_contactos()]),
        ("Animales", ["Caravana / grupo", "Especie", "Categoría", "Cabezas", "Raza", "Rodeo", "Nacimiento",
                      "Estado reproductivo", "Parto probable", "Partos", "Madre", "Situación", "Observaciones"],
         [[a["caravana"], a["especie"], a["categoria"], a["cantidad"], a["raza"], a["rodeo"],
           a["fecha_nacimiento"], _etiqueta(E, a["estado_reproductivo"]) if a["reproductiva"] else "",
           a["fecha_probable_parto"], a["partos"], a["madre_caravana"] or "",
           _etiqueta(opciones.ESTADOS_ANIMAL, a["estado"]), a["observaciones"]] for a in animales]),
        ("Eventos de animales", ["Fecha", "Caravana", "Especie", "Evento", "Resultado", "Crías machos", "Crías hembras", "Detalle"],
         [[e["fecha"], caravana.get(e["animal_id"]), especie.get(e["animal_id"]), _etiqueta(opciones.TIPOS_EVENTO, e["tipo"]),
           _etiqueta(E, e["resultado"]) if e["resultado"] else "", e["crias_machos"] or None,
           e["crias_hembras"] or None, e["detalle"]] for e in sorted(eventos, key=lambda x: x["fecha"], reverse=True)]),
        ("Lotes", ["Lote", "Hectáreas", "Archivado", "Observaciones"],
         [[l["nombre"], l["hectareas"], "Sí" if l["archivado"] else "No", l["observaciones"]]
          for l in lotes_db.listar_lotes(incluir_archivados=True)]),
        ("Cultivos por lote", ["Campaña", "Lote", "Ciclo", "Cultivo", "Variedad / híbrido", "Siembra", "Cosecha",
                               "Hectáreas", "Rinde (qq/ha)", "Observaciones"],
         [[c["campania"], c["lote"], c["ciclo"].capitalize(), c["cultivo"], c["variedad"], c["fecha_siembra"],
           c["fecha_cosecha"], c["hectareas"] or c["hectareas_lote"], c["rinde"], c["observaciones"]]
          for c in lotes_db.listar_lote_cultivos()]),
    ]


@router.get("/exportar/completo")
def exportar_todo(nombre: Optional[str] = "agroapp_completo"):
    return _descarga(crear_excel(hojas_completas()), nombre)
